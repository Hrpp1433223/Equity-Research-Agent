import copy,json
from pathlib import Path
import pytest
from docx import Document
from equity_research_agent.company_reports import analyze,narrative,build,read

ROOT=Path(__file__).resolve().parents[1]
@pytest.mark.parametrize('case',['US_NVDA','HK_00700','SH_600519'])
def test_reference_frame_chart_and_currency(case,tmp_path):
    snapshot=json.loads((ROOT/'tests/fixtures/multi_market'/f'{case}.json').read_text(encoding='utf-8'))
    result=analyze(snapshot);before=copy.deepcopy(result);text=narrative(snapshot,result,tmp_path,False)
    build(snapshot,result,text,tmp_path);doc=Document(tmp_path/'report.docx');manifest=read(tmp_path/'run_manifest.json')
    assert len(doc.tables[0].columns)==2
    assert len(doc.inline_shapes)==1
    assert manifest['chart']['page']==7
    assert manifest['chart']['currency']==snapshot['quote']['currency']
    assert manifest['chart']['values']['base']==pytest.approx(result['model']['base']['quote_value_per_share'])
    assert result==before
    assert not any('Unrelated issuer' in p.text for p in doc.paragraphs)
