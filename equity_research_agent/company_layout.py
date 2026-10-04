"""Company content placed in the retained reference report's page patterns."""
import copy,hashlib,shutil
from pathlib import Path
from docx.shared import Pt
from .reporting import Report
from .charts import dcf_comparison

ROOT=Path(__file__).resolve().parents[1]

def build_reference_report(snapshot,result,text,run):
    from .company_reports import read,write,fmt,pct
    run=Path(run);layout=read(ROOT/'configs/company_report_layout.json');template=None
    working=None
    sec=snapshot['security'];ccy=sec['reporting_currency'];qccy=snapshot['quote']['currency'];fx=snapshot['fx']['quote_units_per_reporting_unit'];model=result['model'];theme=read(ROOT/'configs/report_theme.yaml');last=snapshot['periods'][-1];h=result['history'][-1];m=last['metrics']
    content={'metadata':{'entity':sec['name']},'document_type':'Financial analysis and explicit DCF assumptions','body_header':True,'header':'EQUITY RESEARCH AGENT | PUBLIC-DATA RESEARCH DRAFT','footer':sec['symbol']+' | '+sec['market']+' | '+snapshot['as_of'],'source_caption':'Sources: frozen public financial feeds and matched SEC tags where available. Secondary fields require filing review. Amounts are reporting-currency millions unless stated.'}
    report=Report(content,result.get('inputs') or {},result.get('assumptions') or {},model,result.get('sensitivity') or {},theme,run,template=working);doc=report.doc
    page_roles=[]
    def page(role,title=None,wide=False):
        report.page(report.index+1,wide);page_roles.append({'page':report.index,'role':role,'landscape':wide})
        if title:report.p(doc,title,size=21,line=25,color=theme['cyan'],after=13)
    def para(s,where=None,size=9.5,line=13):return report.p(where if where is not None else doc,s,size=size,line=line,after=9)
    def table(headers,rows,where=None,width=532,label=210,size=8.5,line=14):return report.table(where if where is not None else doc,headers,rows,width=width,label_width=label,size=size,line=line)
    def value(s):return qccy+' '+fmt(model[s]['quote_value_per_share'],2)
    # Reference cover: editorial column and narrow security/valuation sidebar.
    page('cover');left,right=report.frame()
    report.p(left,sec['name']+' | '+sec['market']+' | '+sec['symbol'],size=13,line=17,bold=True,color=theme['purple'])
    report.p(left,'Financial performance and cash value',size=23,line=28,after=13)
    summary='This report connects reported profitability and cash conversion with transparent reinvestment and equity-bridge assumptions.'
    if model:summary+=f" The base case gives {value('base')} per share, compared with the dated close of {qccy} {fmt(snapshot['quote']['close'],2)}. These are conditional DCF values, not an investment rating."
    else:summary+=' Corporate DCF is unavailable for this case; the financial analysis and data limitations remain explicit.'
    para(summary,left,size=10.5,line=15)
    report.heading(left,'Key takeaways',11)
    takeaways=[f"Latest annual revenue {ccy} {fmt(m['revenue']/1e6)} million; EBIT margin {pct(h['ebit_margin'])}; parent net margin {pct(h['parent_net_margin'])}. Fiscal end {last['end']}.",f"Parent ROE {pct(h['parent_roe'])}; CFO / parent profit {fmt(h['cfo_to_parent_income'],2)}x. Scope-matched average parent equity supports the DuPont identity.",f"Dated close {qccy} {fmt(snapshot['quote']['close'],2)} on {snapshot['quote']['date']}. FX {fx:.6f} {qccy} per {ccy}, dated {snapshot['fx']['date']}."]
    if model:takeaways += [f"Bull {value('bull')}, base {value('base')}, bear {value('bear')} per share; base upside / downside {pct(model['base']['upside'])}.",f"WACC {pct(result['assumptions']['wacc'])}; perpetual growth {pct(result['assumptions']['g'])}; stable ROIC {pct(result['assumptions']['terminal_roic'])}; terminal / EV {pct(model['base']['terminal_share'])}. Operating reserve {ccy} {fmt(result['inputs']['bridge']['operating_cash_reserve'])} million is an explicit assumption."]
    for s in takeaways:report.p(left,'• '+s,size=9,line=12.5,after=7)
    para('Financial feeds, operating-profit definitions and forecast assumptions require issuer-filing review. The report does not establish verified fair value. Annual balances and shares are frozen; subsequent distributions, financing and dilution are not automatically bridged.',left,size=9,line=12)
    report.heading(right,'Public-data research draft',10)
    para('Observed figures and calculated values are distinguished from generic scenario assumptions. No assured target price is assigned.',right,size=8.5,line=11)
    rows=[['Ticker',sec['symbol']],['Analysis date',snapshot['as_of']],['Dated close',qccy+' '+fmt(snapshot['quote']['close'],2)]]
    if model:rows += [['Base DCF',value('base')],['DCF / close',pct(model['base']['upside'])],['Bear / bull',fmt(model['bear']['quote_value_per_share'],2)+' / '+fmt(model['bull']['quote_value_per_share'],2)],['WACC / g',pct(result['assumptions']['wacc'])+' / '+pct(result['assumptions']['g'])],['Terminal / EV',pct(model['base']['terminal_share'])]]
    table(['Security / valuation','Value'],rows,right,187,85,8,12)
    report.heading(right,'Snapshot dates',9);para('Annual anchor '+last['end']+'; stock close '+snapshot['quote']['date']+'. Reporting currency '+ccy+'; trading currency '+qccy+'. Current balances are not implied.',right,size=8,line=11)
    report.heading(right,'Bridge proxies',9);para('Cash and recognized short investments less operating reserve, debt, senior claims and book NCI. Unrecognized long-term stakes require separate review.',right,size=8,line=11);report.source(right)
    # Reference cause-analysis typography, immediately after Key Takeaways.
    for s in text['sections']:
        page('cause_analysis',s['title'])
        for p in s['paragraphs']:para(p,size=layout['narrative_font_pt'],line=layout['narrative_line_pt'])
        report.heading(doc,'Evidence and calculation references',10)
        for note in ['[F01] to [F'+f'{len(snapshot["periods"]):02}'+'] identify chronological annual periods in the appendix.','[C01] Calculated DCF values and price comparison. [A01] Explicit assumptions and equity bridge.','[X01] Dated currency conversion. [L01] Source, accounting and model limitations. Full field references are retained in the evidence package.']:para(note,size=8,line=12)
    page('issuer_scope','Issuer scope and reporting comparability');left,right=report.frame()
    for s in [f"The selected security is {sec['name']}, {sec['market']} {sec['symbol']}. The report uses {len(snapshot['periods'])} annual observations; it does not substitute a parent group or another listing.",'SEC-matched periods use actual filing fiscal ends. Other observations retain the vendor period and source status. Publication cutoffs, restatements, mergers and accounting definitions require issuer reconciliation; no uninterrupted economic perimeter is inferred from availability alone.','Parent-attributable profit and average parent equity are matched for ROE. Consolidated operating cash flow / parent profit is a diagnostic comparison, not a scope-matched accounting identity.']:para(s,left)
    report.heading(right,'Issuer identifiers',10);para(sec['name']+'\n'+sec['market']+' '+sec['symbol']+'\nReporting: '+ccy+'\nQuote: '+qccy,right,size=9,line=13)
    report.heading(right,'Version discipline',9);para('Snapshots retain original values, field references, URL provenance and response hashes. Missing data stays unavailable. Source status is not upgraded merely because arithmetic passes.',right,size=8.5,line=12);report.source(doc)
    page('ratios','Profitability cash conversion and DuPont')
    para('Profit margin, cash conversion and balance-sheet returns answer different questions. Parent ROE uses parent-attributable profit and average parent equity; the three-factor identity uses matching balances throughout. Missing opening balances remain unavailable.')
    metrics=[('growth','Revenue growth'),('ebit_margin','EBIT margin'),('parent_net_margin','Parent net margin'),('cfo_to_parent_income','CFO / parent profit'),('parent_roe','Parent ROE, average equity'),('asset_turnover','Asset turnover'),('assets_to_parent_equity','Assets / parent equity')]
    table(['Metric']+[x['end'] for x in result['history']],[[label]+[fmt(x[k],2)+'x' if k in {'cfo_to_parent_income','asset_turnover','assets_to_parent_equity'} and x[k] is not None else pct(x[k]) for x in result['history']] for k,label in metrics],label=205,size=8,line=15)
    table(['DuPont reconciliation']+[x['end'] for x in result['history']],[[label]+[pct(x[k]) if k!='dupont_error' else ('Not available' if x[k] is None else f'{x[k]:.2e}') for x in result['history']] for k,label in [('parent_roe','Direct parent ROE'),('dupont_product','Three-factor product'),('dupont_error','Absolute identity error')]],label=205,size=8,line=15);report.source(doc)
    chart_manifest=None
    if model:
        base=model['base'];pars=result['driver_parameters'];periods=base['periods']
        page('forecast','Five year operating forecast')
        para('Forecasts are explicit assumptions, not management guidance. The model projects the issuer in aggregate; there is no verified segment, product or order forecast in this packet. The first annual FCFF is prorated to the days remaining after the analysis date.')
        table(['Base drivers / forecast']+[p['year'] for p in periods],[[label]+[pct(p[k]) if k in {'margin','tax_rate'} else fmt(p[k]) for p in periods] for k,label in [('revenue','Revenue'),('margin','EBIT margin'),('ebit','EBIT'),('tax_rate','Operating tax rate'),('da','D&A'),('capex','Capex'),('nwc','Trade NWC proxy'),('annual_fcff','Full year FCFF')]],label=150,size=8,line=15)
        para('Initial growth '+pct(pars['revenue_growth'])+' fades toward perpetual growth. Historical D&A / revenue '+pct(pars['da_ratio'])+', capex / revenue '+pct(pars['capex_ratio'])+' and trade NWC / revenue '+pct(pars['nwc_ratio'])+' are mechanical drivers unless explicitly overridden. Vendor EBIT requires accounting reconciliation before interpretation as core operating profit.',size=9,line=12);report.source(doc)
        page('dcf_comparison','DCF comparison | Three operating scenarios');left,right=report.frame()
        report.heading(left,'BASE DCF '+value('base'),11);para('FCFF / WACC as of '+snapshot['as_of']+'. Five explicit years and normalized terminal reinvestment.',left,size=9,line=12)
        charts=run/'charts';charts.mkdir(exist_ok=True);quoted=copy.deepcopy(model)
        for v in quoted.values():v['value_per_share']=v['quote_value_per_share']
        chart=charts/'dcf_comparison.png';dcf_comparison(quoted,snapshot['quote']['close'],chart,currency=qccy,as_of=snapshot['as_of'])
        p=left.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4);p.add_run().add_picture(str(chart),width=Pt(325))
        para('Connections compare values with the dated close. No price path or scenario probability is assigned.',left,size=7.5,line=10)
        report.heading(right,'Shared valuation inputs',10);table(['Assumption','Value'],[['WACC',pct(base['wacc'])],['Perpetual growth',pct(base['g'])],['Stable ROIC',pct(result['assumptions']['terminal_roic'])],['Stable tax',pct(result['assumptions']['terminal_tax_rate'])],['Base stable margin',pct(base['terminal_margin'])],['TV / EV',pct(base['terminal_share'])]],right,187,125,8,12)
        para('WACC and stable ROIC are explicit assumptions. Terminal share measures sensitivity; it does not validate forecast accuracy.',right,size=8,line=11)
        report.heading(doc,'Operating cases',11);columns=doc.add_table(rows=1,cols=3);columns.autofit=False
        from .assumption_advisor import DEFAULT_SHIFTS
        shifts=pars.get('scenario_shifts',DEFAULT_SHIFTS)
        descriptions={'base':'Recent growth fades toward perpetual growth. Margin, D&A and reinvestment ratios follow the disclosed base assumptions.'}
        for s in ['bull','bear']:descriptions[s]='Shifts from base: growth '+f"{shifts[s+'_growth_shift']*100:+.1f}"+' pp; margin '+f"{shifts[s+'_margin_shift']*100:+.1f}"+' pp; capex ratio '+f"{shifts[s+'_capex_shift']*100:+.1f}"+' pp; NWC ratio '+f"{shifts[s+'_nwc_shift']*100:+.1f}"+' pp.'
        for i,s in enumerate(['bull','base','bear']):
            cell=columns.cell(0,i);cell.width=Pt(177);report.heading(cell,s.upper()+' '+value(s),10);para('DCF / close '+pct(model[s]['upside']),cell,size=8.5,line=11);para(descriptions[s],cell,size=9,line=12);para(model[s]['periods'][-1]['year']+' revenue '+fmt(model[s]['periods'][-1]['revenue'])+'m; EBIT margin '+pct(model[s]['periods'][-1]['margin'])+'. Stable margin '+pct(model[s]['terminal_margin'])+'.',cell,size=8,line=11)
        report.source(doc);chart_manifest={'file':'charts/dcf_comparison.png','sha256':hashlib.sha256(chart.read_bytes()).hexdigest(),'currency':qccy,'reference_close':snapshot['quote']['close'],'values':{s:v['quote_value_per_share'] for s,v in model.items()},'page':report.index}
        page('dcf_detail','DCF calculation detail | '+ccy+' million',True)
        rows=[]
        for s,v in model.items():
            for p in v['periods']:rows.append([s.upper(),p['year']]+[fmt(p[k]) for k in ['revenue','ebit','nopat','da','capex','delta_nwc','annual_fcff']]+[f'{p["portion"]:.4f}',fmt(p['fcff']),f'{p["discount_factor"]:.5f}',fmt(p['pv'])])
        table(['Case','Year','Revenue','EBIT','NOPAT','D&A','Capex','Change NWC','Annual FCFF','Fraction','FCFF','DF','PV'],rows,width=712,label=44,size=7,line=14)
        table(['Terminal / value','Bull','Base','Bear'],[[label]+[fmt(model[s][k]) for s in ['bull','base','bear']] for k,label in [('terminal_fcff','Stable FCFF'),('terminal_value','Terminal value'),('terminal_pv','PV of terminal'),('explicit_pv','PV of explicit'),('ev','Operating EV'),('equity','Equity value')]]+[[qccy+' / share']+[fmt(model[s]['quote_value_per_share'],2) for s in ['bull','base','bear']]],width=712,label=240,size=8,line=14)
        para('Annual FCFF = NOPAT + D&A - capex - change in trade NWC. The first fraction and actual/365 discount factor are shown explicitly. Stable FCFF = normalized NOPAT x (1 - growth / ROIC). Amounts stay in '+ccy+'; only per-share outputs are translated.',size=8,line=11)
        page('equity_bridge','Equity bridge and sensitivity');left,right=report.frame();bridge=result['inputs']['bridge'];sens=result['sensitivity']
        rows=[['Operating EV',fmt(base['ev'])]]+[[k.replace('_',' ').title(),fmt(v,6 if k=='shares' else 1)] for k,v in bridge.items()]+[['Equity value',fmt(base['equity'])],[qccy+' / share',fmt(base['quote_value_per_share'],2)]]
        table(['Base bridge | '+ccy+' mm','Amount'],rows,left,325,210,8,13)
        table(['WACC / g']+[pct(g) for g in sens['g']],[[pct(w)]+['Invalid' if v is None else fmt(v*fx,2) for v in row] for w,row in zip(sens['wacc'],sens['values'])],right,187,39,7,14)
        para('Sensitivity in '+qccy+'/share. Centre equals the base model. Other operating and bridge assumptions are held fixed.',right,size=8,line=11)
        for s in [f"Reporting currency {ccy}; quoted currency {qccy}. Conversion uses {fx:.6f} {qccy} per {ccy}, dated {snapshot['fx']['date']}.",'Cash and recognized short investments enter the bridge after the operating reserve. Debt, senior claims and NCI are deducted once; additional non-operating asset recognition is a separate explicit policy.','NCI is a book-value proxy. Annual cash, debt and shares are frozen unless overridden. Subsequent distributions, buybacks, financing, dilution and changes in investment fair value are not automatically bridged.']:para(s,size=9,line=12)
        report.source(doc)
    page('history','Historical financial appendix',True)
    fields=[('revenue','Revenue'),('ebit','Reported operating EBIT'),('parent_income','Parent net income'),('cfo','Cash from operations'),('capex','Capital expenditure'),('da','Depreciation / amortization'),('assets','Total assets'),('equity','Parent equity'),('cash','Cash equivalents'),('investments','Recognized short investments'),('debt','Debt / leases'),('nci','NCI book balance'),('shares','Ordinary shares million')]
    table([ccy+' million']+[p['end'] for p in snapshot['periods']],[[label]+[fmt(p['metrics'].get(k)/1e6 if p['metrics'].get(k) is not None else None) for p in snapshot['periods']] for k,label in fields],width=712,label=205,size=9,line=19)
    para('Periods retain their fiscal dates and source definitions. Historical share units may retain pre-split observations and are not used to infer per-share growth. Unavailable fields are missing observations, not factual zero balances.',size=9,line=13);report.source(doc)
    page('method','Method evidence and limitations')
    para('FCFF discounts operating cash flows at WACC; financial claims enter the equity bridge. Terminal cash flow normalizes reinvestment using growth / stable ROIC. WACC must exceed perpetual growth. The chart compares scenario values with a dated close, not a future price path.')
    table(['Source ID','Kind / host','Retrieved'],[['S'+str(i),s['source_kind']+' / '+s['url'].split('/')[2],s['retrieved_utc'][:10]] for i,s in enumerate(snapshot['sources'],1)],label=85,size=8,line=14)
    for s in result['warnings']:para(s,size=8,line=11)
    source_ids={s['url']:'S'+str(i) for i,s in enumerate(snapshot['sources'],1)};locators=[]
    for i,p in enumerate(snapshot['periods'],1):
        for k,v in p['metrics'].items():
            ref=p['refs'].get(k,{});tag=ref.get('tag') or ref.get('field') or ('Component subtotal' if ref.get('components') else 'See evidence');tag=tag if len(tag)<=34 else tag[:31]+'...'
            locators.append([f'F{i:02}.'+p['end']+'.'+k,fmt(v,2),'shares' if k=='shares' else ccy,source_ids.get(ref.get('url'),'See JSON'),tag])
    count=layout['source_locator_rows_per_page']
    for offset in range(0,len(locators),count):
        page('source_locators','Financial source locators | '+str(offset+1)+'-'+str(min(offset+count,len(locators))),True)
        table(['Fact ID / fiscal end','Raw value','Unit','Source ID','Field locator'],locators[offset:offset+count],width=712,label=235,size=7,line=9.4)
        para('Field locators may be shortened for display. Full tags, source URLs, accession/filing dates where matched, raw responses, hashes and component definitions remain in company_snapshot.json and data/. SEC matching does not replace accounting-scope review; secondary feeds are not independently filing-verified.',size=7.5,line=10)
    if (run/'assumption_discussion.json').exists():
        discussion=read(run/'assumption_discussion.json');adopted=discussion['adopted']
        for offset in range(0,max(1,len(adopted)),4):
            page('assumption_discussion','Adopted DCF assumptions and reasoning')
            para('These proposals were selected by the user and still equal the final input values. Manual changes remain user inputs. Discussion does not independently verify sources or make assumptions facts.',size=9,line=13)
            if not adopted:para('No DeepSeek suggestion was applied unchanged to the final DCF. The discussion remains in the evidence package.')
            for x in adopted[offset:offset+4]:
                shown=lambda v:fmt(v,2)+' million' if x['parameter']=='reserve' else f'{v*100:.2f}'+(' pp' if x['parameter'].endswith('_shift') else '%')
                report.heading(doc,x['parameter'].replace('_',' ').title()+' '+shown(x['value']),10)
                para('Range '+shown(x['low'])+' to '+shown(x['high'])+'; basis '+x['basis'].replace('_',' ')+'; sources '+(', '.join(x['source_ids']) or 'no retrieved evidence')+'.',size=8,line=11)
                para(x['rationale'][:320],size=9,line=13);para('Counterargument: '+x['counterargument'][:240],size=8.5,line=12)
            para('Full discussions, parameter overrides, prompt, source URLs, retrieval timestamps and raw documents are retained in assumption_discussion.json and discussion_research/.',size=8,line=11)
        page('discussion_sources','Discussion research sources and limitations')
        sources=discussion['research']['sources']
        table(['ID','Public source host','Published','Evidence status'],[[s['id'],s['url'].split('/')[2],str(s.get('published_at') or 'Unknown')[:25] if not s['status'].startswith('official_index') else 'Unknown',s['status'].replace('_',' ')] for s in sources],label=35,size=8,line=14)
        para('Sources were retrieved for discussion and have not been independently reconciled with issuer filings. Search snippets are discovery evidence. Unknown publication dates prevent a reliable publication cutoff check. Raw source titles and URLs remain in the evidence package.',size=9,line=13)
        para('The five-year model uses a growth fade and constant margin and reinvestment ratios with disclosed scenario shifts. It does not support arbitrary yearly or segment forecasts. All suggested value changes are calculated by the local financial engine, not by DeepSeek.',size=9,line=13)
    for section in doc.sections[1:]:section.header.is_linked_to_previous=True
    doc.core_properties.subject='Public financial analysis with explicit assumptions using the retained reference layout'
    doc.save(run/'report.docx')
    manifest={'case_id':'multi_market','as_of':snapshot['as_of'],'expected_pages':report.index,'docx_sha256':hashlib.sha256((run/'report.docx').read_bytes()).hexdigest(),'output_language':'en','snapshot_status':snapshot['status'],'llm_used':(run/'provider_metadata.json').exists(),'template_version':layout['version'],'template_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'page_roles':page_roles,'chart':chart_manifest}
    write(run/'run_manifest.json',manifest)
