"""Editable page recipes, independent from company content and financial truth."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def clean(s):return s.replace('��',' - ').replace('�','-').replace('—','-').replace('–','-').replace('’',"'").replace('“','"').replace('”','"').replace('\u00a0',' ')

def resolve(s,model,assumptions=None):
    for scenario,val in model.items():
        prefix=val.get('currency')+' ' if val.get('currency') in ('CNY','HKD') else '$'
        s=s.replace('{{metric:'+scenario+'.value_per_share}}',f"{prefix}{val['value_per_share']:.2f}")
        s=s.replace('{{metric:'+scenario+'.upside}}',f"{val['upside']:+.1%}")
    for k,v in (assumptions or {}).items():
        if isinstance(v,(int,float)):s=s.replace('{{assumption:'+k+'}}',f'{v:.1%}')
    return clean(s)

class Report:
    def __init__(self,content,inputs,assumptions,model,sens,theme,run,template=None):
        self.c=content;self.inputs=inputs;self.a=assumptions;self.m=model;self.sens=sens;self.t=theme;self.run=Path(run);self.doc=Document(template) if template else Document();self.index=0
        if template:
            # Retain the supplied styles/theme/settings while replacing company slots.
            body=self.doc._element.body
            for element in list(body):
                if element.tag!=qn('w:sectPr'):body.remove(element)
            for element in list(body.sectPr):
                if element.tag in (qn('w:headerReference'),qn('w:footerReference')):body.sectPr.remove(element)
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            for key,relation in list(self.doc.part.rels.items()):
                if relation.reltype in (RT.IMAGE,RT.HEADER,RT.FOOTER):self.doc.part.drop_rel(key)
        style=self.doc.styles['Normal']; style.font.name=theme['font'];style.font.size=Pt(theme['body_size']);style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=Pt(theme['body_line'])
        for name in ['Title','Heading 1','Heading 2','Heading 3']:
            st=self.doc.styles[name];st.font.name=theme['font'];st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(6)
        self.doc.core_properties.author='Equity Research Agent';self.doc.core_properties.title=content['metadata']['entity']+' '+content.get('document_type','Reference reproduction and DCF adaptation')
        self.doc.core_properties.subject='Local reference reproduction; no institutional endorsement'

    def p(self,where,s='',size=None,bold=False,color=None,after=6,line=None):
        p=where.add_paragraph();p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(line or ((size or self.t['body_size'])*1.35));p.paragraph_format.widow_control=False
        r=p.add_run(resolve(s,self.m,self.a));r.bold=bold;r.font.name=self.t['font'];r.font.size=Pt(size or self.t['body_size'])
        if size and size>=20:r.font.name='Segoe UI Light'
        if color:r.font.color.rgb=RGBColor.from_string(color)
        return p

    def page(self,number,landscape=False):
        sec=self.doc.sections[0] if self.index==0 else self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        self.index+=1;sec.page_width=Pt(792 if landscape else 612);sec.page_height=Pt(612 if landscape else 792);sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.left_margin=sec.right_margin=Pt(40);sec.top_margin=Pt(94 if landscape else 80);sec.bottom_margin=Pt(36);sec.header_distance=Pt(30);sec.footer_distance=Pt(15)
        sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
        h=sec.header.paragraphs[0];h.paragraph_format.line_spacing=Pt(14);r=h.add_run(self.c.get('header','EQUITY RESEARCH AGENT  |  REFERENCE REPRODUCTION'));r.font.name=self.t['font'];r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string(self.t['blue'])
        f=sec.footer.paragraphs[0];f.paragraph_format.line_spacing=Pt(8);r=f.add_run(self.c.get('footer','Reference reproduction · DCF adaptation  |  26 Jan 2026'));r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string('607D90')
        f.add_run(' '*8);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
        # Empty section-break paragraphs must not consume body space.
        if self.doc.paragraphs:
            p=self.doc.paragraphs[-1];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1)
            for r in p.runs:r.font.size=Pt(1)
        if self.c.get('body_header',False):
            # Explicitly paged case reports render a stable masthead in body flow.
            # Word COM inconsistently displays repeated section-header text.
            h.text='';sec.top_margin=Pt(44 if landscape else 30)
            self.p(self.doc,self.c.get('header','EQUITY RESEARCH AGENT'),size=10,line=14,color=self.t['blue'],after=36)
        return self.doc

    def table(self,where,headers,rows,width=532,label_width=130,size=7,line=8.4,groups=None):
        n=len(headers); t=where.add_table(rows=1,cols=n);t.autofit=False
        widths=[label_width]+[(width-label_width)/(n-1)]*(n-1)
        for col,w in zip(t.columns,widths):col.width=Pt(w)
        tPr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','left','right','insideH','insideV']:
            b=OxmlElement('w:'+edge);b.set(qn('w:val'),'nil');borders.append(b)
        tPr.append(borders)
        def fillrow(cells,values,header=False,bold=False):
            for i,(cell,s) in enumerate(zip(cells,values)):
                cell.width=Pt(widths[i]); tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'0' if side in ['top','bottom'] else '15');e.set(qn('w:type'),'dxa');mar.append(e)
                tcPr.append(mar);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(line);p.paragraph_format.widow_control=False
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
                rendered=clean(str(s)).replace('ModelWare','Reference adjusted')
                if width<=340:rendered=rendered.replace('Consensus','Cons.').replace('Effective Tax Rate','Tax rate').replace('Total Revenue','Revenue').replace('Pre-tax Profits','Pretax profit')
                r=p.add_run(rendered);r.font.name='Arial Narrow' if size<=7.5 else self.t['font'];r.font.size=Pt(size);r.bold=header or bold
                if i>0:tcPr.append(OxmlElement('w:noWrap'))
                if header:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),self.t['blue']);tcPr.append(sh);r.font.color.rgb=RGBColor(255,255,255)
                elif i>0 and re.search(r'\d',str(s)):r.font.color.rgb=RGBColor.from_string('082C4A')
            trPr=cells[0]._tc.getparent().get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
            if header:trPr.append(OxmlElement('w:tblHeader'))
        fillrow(t.rows[0].cells,headers,True)
        if groups:
            first=t.add_row(); t._tbl.remove(first._tr);t._tbl.insert(2,first._tr)
            for start,end,label in groups:
                cell=first.cells[start].merge(first.cells[end]);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(label);r.font.name=self.t['font'];r.font.size=Pt(size);r.bold=True;r.font.color.rgb=RGBColor.from_string(self.t['blue'])
            first._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows:
            vals=row['cells'] if isinstance(row,dict) else row
            fillrow(t.add_row().cells,vals,bold=row.get('bold',False) if isinstance(row,dict) else False)
        return t

    def frame(self):
        t=self.doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Pt(333);t.columns[1].width=Pt(199)
        for i,cell in enumerate(t.rows[0].cells):
            cell.width=Pt([333,199][i]);pr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,twips in [('top',0),('bottom',0),('left',0 if i==0 else 180),('right',160 if i==0 else 0)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(twips));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run('').font.size=Pt(1)
        borders=OxmlElement('w:tcBorders');b=OxmlElement('w:right');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'3');b.set(qn('w:color'),'8F969C');borders.append(b);t.cell(0,0)._tc.get_or_add_tcPr().append(borders)
        return t.cell(0,0),t.cell(0,1)

    def fig(self,where,ex,width=325):
        p=where.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(self.run/'charts'/f'exhibit_{ex}.png'),width=Pt(width))
    def source(self,where,extra=''):
        self.p(where,self.c.get('source_caption','Source: reference report, 26 Jan 2026 (supplied PDF).')+' '+extra,size=7,line=8.2,color='666666',after=7)
    def body(self,where,page,group=0):
        for para in self.c['pages'][str(page)]['paragraphs'][group]:
            text=para['text'];self.p(where,text,bold=False,size=9.2,line=12.4 if page==6 else 13.5,after=7)
    def heading(self,where,s,size=11):self.p(where,s,size=size,bold=True,after=5,line=size*1.25)

